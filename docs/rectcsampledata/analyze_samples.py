import openpyxl
import sys
from pdfminer.high_level import extract_text

BASE = "/Users/sujoymukherjee/code/ctc/CTC Final/docs/rectcsampledata"

def analyze_excel(filepath, title):
    print("=" * 80)
    print(title)
    print("=" * 80)
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for shname in wb.sheetnames:
            ws = wb[shname]
            print(f"\n--- Sheet: '{shname}' (dims: {ws.dimensions}, max_row:{ws.max_row}, max_col:{ws.max_column}) ---")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if any(c is not None for c in row):
                    print(f"  Row {i+1:3d}: {list(row)}")
                if i > 80:
                    print("  [...further rows truncated...]")
                    break
    except Exception as e:
        print(f"ERROR: {e}")

def analyze_pdf(filepath, title):
    print("=" * 80)
    print(title)
    print("=" * 80)
    try:
        text = extract_text(filepath)
        print(text[:4000])
        if len(text) > 4000:
            print(f"\n[...{len(text)-4000} chars truncated...]")
    except Exception as e:
        print(f"ERROR: {e}")

analyze_excel(f"{BASE}/CTC DESPATCH REGISTER Sample.xlsx", "DISPATCH REGISTER SAMPLE")
analyze_excel(f"{BASE}/CTC BILL NOTEBOOK Sample.xlsx", "BILL NOTEBOOK SAMPLE")

pdfs = [
    ("Hire Memo - Empty.pdf", "HIRE MEMO - EMPTY (TEMPLATE)"),
    ("Hire Memo - Filled.pdf", "HIRE MEMO - FILLED"),
    ("Invoice - Empty.pdf", "INVOICE - EMPTY (TEMPLATE)"),
    ("Invoice Filled - With Annexure.pdf", "INVOICE FILLED WITH ANNEXURE"),
    ("Invoice Filled - without annexure.pdf", "INVOICE FILLED WITHOUT ANNEXURE"),
    ("Lorry Receipt - Filled.pdf", "LORRY RECEIPT - FILLED"),
    ("Lorry Reciept (LR) - Empty.pdf", "LORRY RECEIPT - EMPTY (TEMPLATE)"),
]
for fname, title in pdfs:
    analyze_pdf(f"{BASE}/{fname}", title)
